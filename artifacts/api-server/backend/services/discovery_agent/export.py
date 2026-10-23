"""Excel export of discovered events and their organizers — the download on
the admin "Events found" page.

Every value is written as plain text/number/date, never as a formula:
scraped titles like `=HYPERLINK(...)` would otherwise run when the file is
opened (spreadsheet formula injection).
"""

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from backend.models import DiscoveredEvent, DiscoveredOrganizer

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_EVENT_COLUMNS = (
    ("Event name", lambda e: e.name),
    ("Event date", lambda e: e.event_date),
    ("Date as written", lambda e: e.date),
    ("Start time", lambda e: e.start_time),
    ("End time", lambda e: e.end_time),
    ("Venue", lambda e: e.venue),
    ("City", lambda e: e.city),
    ("State", lambda e: e.state),
    ("Format", lambda e: e.event_format),
    ("Category", lambda e: e.category),
    ("Type", lambda e: e.subcategory),
    ("Event link", lambda e: e.url),
    ("Ticket link", lambda e: e.ticket_url),
    ("Verification", lambda e: e.verification_status),
    ("Confidence", lambda e: e.confidence_score),
    ("Source website", lambda e: e.source_website),
    ("Organizer name", lambda e: e.organizer.name if e.organizer else None),
    ("Organizer email", lambda e: e.organizer.email if e.organizer else None),
    ("Found on", lambda e: e.discovered_at),
)

_ORGANIZER_COLUMNS = (
    ("Organizer name", lambda o: o.name),
    ("Email", lambda o: o.email),
    ("Phone", lambda o: o.phone),
    ("Website", lambda o: o.website),
    ("LinkedIn", lambda o: o.linkedin),
    ("Type", lambda o: o.organizer_type),
    ("Confidence", lambda o: o.confidence_score),
    ("Events in this file", None),
)


def build_events_workbook(events: list[DiscoveredEvent]) -> bytes:
    book = Workbook()
    events_sheet = book.active
    events_sheet.title = "Events"
    _write_sheet(events_sheet, [name for name, _ in _EVENT_COLUMNS], ([getter(e) for _, getter in _EVENT_COLUMNS] for e in events))

    organizers: dict[int, tuple[DiscoveredOrganizer, int]] = {}
    for event in events:
        if event.organizer is not None:
            organizer, count = organizers.get(event.organizer.id, (event.organizer, 0))
            organizers[organizer.id] = (organizer, count + 1)
    organizer_rows = (
        [getter(o) for _, getter in _ORGANIZER_COLUMNS if getter is not None] + [count]
        for o, count in organizers.values()
    )
    _write_sheet(book.create_sheet("Organizers"), [name for name, _ in _ORGANIZER_COLUMNS], organizer_rows)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _write_sheet(sheet, header: list[str], rows) -> None:
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_plain(value) for value in row])
        for cell in sheet[sheet.max_row]:
            _never_formula(cell)
    sheet.freeze_panes = "A2"
    for index, name in enumerate(header, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(12, min(48, len(name) + 6))
    sheet.column_dimensions["A"].width = 48


def _plain(value):
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)  # Excel has no time zones
    if isinstance(value, (date, int, float)) or value is None:
        return value
    return str(value)


def _never_formula(cell: Cell) -> None:
    if cell.data_type == "f":
        cell.data_type = "s"
