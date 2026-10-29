"""Event Discovery automation (no AI): DuckDuckGo web search, page reading
with a best-effort text reader for pages without structured data, and the
date-range Excel export of what was found."""

import io
import json
from datetime import date, timedelta

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.db import get_db
from backend.dependencies import get_current_user
from backend.models import AgentRun, DiscoveredEvent, DiscoveredOrganizer, User
from backend.routes import discovery as discovery_routes
from backend.services.discovery_agent import orchestrator, page_extractor, text_extractor, validation, web_search
from backend.services.discovery_agent.types import CandidateSource, DiscoveryConfig, OrganizerCandidate

NEXT_YEAR = date.today().year + 1


# --- Web search ------------------------------------------------------------------

_DDG_RESULTS = """<html><body>
<div class="result results_links web-result result--ad"><a class="result__a" href="https://duckduckgo.com/y.js?ad_domain=x.com&u3=1">Ad</a></div>
<div class="result results_links web-result"><a class="result__a" href="https://allevents.in/mumbai/business">Business Events in Mumbai</a></div>
<div class="result results_links web-result"><a class="result__a" href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fwww.bengalurutechsummit.com%2F&rut=abc">Bengaluru Tech Summit</a></div>
<div class="result results_links web-result"><a class="result__a" href="https://allevents.in/mumbai/business">Business Events in Mumbai (again)</a></div>
</body></html>"""


class _Response:
    def __init__(self, status_code, text):
        self.status_code = status_code
        self.text = text


# The real parser, captured at import — conftest's no_live_web_search stubs
# every provider's search() for each test; these tests fake only the HTTP call.
from backend.services.discovery_agent.search_providers import web_fallback  # noqa: E402

_REAL_DDG_SEARCH = web_fallback.WebFallbackProvider.search


def _ddg(monkeypatch, status, body):
    """The keyless DuckDuckGo provider (search_providers/web_fallback.py) answering with `body`."""
    monkeypatch.setattr(web_fallback.WebFallbackProvider, "search", _REAL_DDG_SEARCH)
    monkeypatch.setattr(web_fallback.httpx, "post", lambda *args, **kwargs: _Response(status, body))
    monkeypatch.setattr(web_fallback.WebFallbackProvider, "_wait_turn", lambda self: None)
    return web_fallback.WebFallbackProvider()


def test_search_results_are_parsed_unwrapped_and_ads_skipped(monkeypatch):
    provider = _ddg(monkeypatch, 200, _DDG_RESULTS)
    assert [r.url for r in provider.search("business events Mumbai", max_results=10)] == [
        "https://allevents.in/mumbai/business",
        "https://www.bengalurutechsummit.com/",
    ]


def test_search_returns_capped_results(monkeypatch):
    provider = _ddg(monkeypatch, 200, _DDG_RESULTS)
    assert [r.url for r in provider.search("business events Mumbai", max_results=1)] == ["https://allevents.in/mumbai/business"]


@pytest.mark.parametrize("status, body", [(202, "<html>anomaly</html>"), (200, '<div class="anomaly-modal">Unfortunately, bots use DuckDuckGo too.</div>'), (429, "")])
def test_blocked_search_raises(monkeypatch, status, body):
    provider = _ddg(monkeypatch, status, body)
    with pytest.raises(RuntimeError):
        provider.search("anything", max_results=5)


# --- Best-effort text reader -------------------------------------------------------

_PLAIN_EVENT_PAGE = f"""<html><head><title>Bengaluru Healthcare Summit {NEXT_YEAR} | BTS</title>
<meta property="og:title" content="Bengaluru Healthcare Summit {NEXT_YEAR}"></head><body>
<h1>Bengaluru Healthcare Summit {NEXT_YEAR}</h1>
<p>Held on 1 January 2020 for the first time.</p>
<p>Dates: 18-20 November {NEXT_YEAR}</p>
<p>Venue: Bangalore International Exhibition Centre (BIEC), Tumkur Road</p>
</body></html>"""


def _source(url="https://www.bengalurutechsummit.com/"):
    return CandidateSource(url=url, source_type="event_listing", source_website="www.bengalurutechsummit.com")


def test_best_effort_reads_title_first_upcoming_date_and_venue():
    event = text_extractor.extract_best_effort(_source(), _PLAIN_EVENT_PAGE, cities=["Bengaluru", "Mumbai"])

    assert event is not None
    assert event.name == f"Bengaluru Healthcare Summit {NEXT_YEAR}"
    assert event.date == f"{NEXT_YEAR}-11-18"  # the past 2020 date is skipped
    assert event.venue == "Bangalore International Exhibition Centre (BIEC), Tumkur Road"
    assert event.city == "Bengaluru"
    assert event.best_effort is True
    assert event.url == "https://www.bengalurutechsummit.com/"


@pytest.mark.parametrize("text, expected", [
    ("November 18, {y}", "{y}-11-18"),
    ("Nov 18th - 20th, {y}", "{y}-11-18"),
    ("18th Nov {y}", "{y}-11-18"),
    ("{y}-11-18", "{y}-11-18"),
    ("Sept 3 {y}", "{y}-09-03"),
])
def test_best_effort_date_formats(text, expected):
    html = f"<html><h1>Some Summit</h1><p>{text.format(y=NEXT_YEAR)}</p></html>"
    event = text_extractor.extract_best_effort(_source(), html, cities=[])

    assert event is not None and event.date == expected.format(y=NEXT_YEAR)


def test_best_effort_needs_a_real_upcoming_date():
    html = "<html><h1>Some Summit</h1><p>Last held on 12 March 2021.</p></html>"

    assert text_extractor.extract_best_effort(_source(), html, cities=[]) is None


def test_best_effort_ignores_listing_pages():
    html = f"<html><title>Upcoming Business Events in Mumbai</title><h1>Business events in Mumbai</h1><p>12 December {NEXT_YEAR}</p></html>"

    assert text_extractor.extract_best_effort(_source(), html, cities=["Mumbai"]) is None


def test_best_effort_events_always_need_review_even_with_a_low_threshold():
    event = text_extractor.extract_best_effort(_source(), _PLAIN_EVENT_PAGE, cities=["Bengaluru"])

    status, notes, _, _ = validation.validate_event(event, confidence_threshold=0)

    assert status == "NEEDS_REVIEW"
    assert any("fixed rules" in note for note in notes)


# --- Reading any page ---------------------------------------------------------------


def _serve(monkeypatch, pages):
    monkeypatch.setattr(page_extractor, "fetch", lambda url, render_js=False: pages.get(url))
    monkeypatch.setattr(page_extractor, "browser_fetch_enabled", lambda: False)


def _ld(node) -> str:
    return f'<script type="application/ld+json">{json.dumps(node)}</script>'


def test_single_event_page_without_its_own_url_uses_the_page_url(monkeypatch):
    url = "https://www.nasscom.in/ntc/"
    _serve(monkeypatch, {url: "<html><head>" + _ld({"@type": "Event", "name": "NASSCOM Technology Conference", "startDate": f"{NEXT_YEAR}-02-10"}) + "</head></html>"})

    events = page_extractor.extract_events_from_page(url, max_events=10, cities=[])

    assert [e.name for e in events] == ["NASSCOM Technology Conference"]
    assert events[0].url == url
    assert events[0].best_effort is False


def test_page_without_structured_data_or_event_links_falls_back_to_best_effort(monkeypatch):
    url = "https://www.bengalurutechsummit.com/"
    _serve(monkeypatch, {url: _PLAIN_EVENT_PAGE})

    events = page_extractor.extract_events_from_page(url, max_events=10, cities=["Bengaluru"])

    assert len(events) == 1
    assert events[0].best_effort is True


# --- Orchestrator: no AI anywhere -------------------------------------------------------


def _config(**overrides) -> DiscoveryConfig:
    values = dict(cities=["Bengaluru"], states=[], categories=["Healthcare"], search_depth=2, sources_per_query=3,
                  daily_limit=200, confidence_threshold=60, listing_urls=[], events_per_listing=10, require_organizer_email=False)
    values.update(overrides)
    return DiscoveryConfig(**values)


def _no_organizer(monkeypatch):
    monkeypatch.setattr(orchestrator.organizer_identifier, "identify_organizer", lambda candidate: OrganizerCandidate(found=False))


def test_run_searches_the_web_and_stores_events_without_ai(db_session, monkeypatch):
    url = "https://www.bengalurutechsummit.com/"
    _serve(monkeypatch, {url: _PLAIN_EVENT_PAGE})
    _no_organizer(monkeypatch)
    queries: list[str] = []
    monkeypatch.setattr(orchestrator.web_search, "search", lambda query, max_results=5, **kwargs: queries.append(query) or [url])

    summary = orchestrator.run_discovery_agent(db_session, _config())

    assert len(queries) == 2
    assert summary.new_events == 1  # the second query found the same page again — read once per run
    event = db_session.query(DiscoveredEvent).one()
    assert event.verification_status == "NEEDS_REVIEW"
    run = db_session.query(AgentRun).filter(AgentRun.agent_type == "discovery").one()
    assert run.status == "completed"
    assert run.error is None


def test_blocked_search_keeps_listing_results_and_reports_it(db_session, monkeypatch):
    # The run is scoped to Healthcare, so the listed event has to be a healthcare one
    # ("Startup Mixer" reads as Finance and would be set aside as off-sector).
    listing = "https://allevents.in/bangalore"
    event_node = {"@type": "Event", "name": "Hospital Doctors Mixer", "url": "https://allevents.in/bangalore/hospital-doctors-mixer/1", "startDate": f"{NEXT_YEAR}-03-01",
                  "location": {"name": "WeWork Galaxy", "address": {"addressLocality": "Bengaluru"}}}
    _serve(monkeypatch, {listing: "<html><head>" + _ld(event_node) + "</head></html>"})
    _no_organizer(monkeypatch)

    def blocked(query, max_results=5, **kwargs):
        raise web_search.SearchBlocked("DuckDuckGo is asking for a captcha")

    monkeypatch.setattr(orchestrator.web_search, "search", blocked)

    summary = orchestrator.run_discovery_agent(db_session, _config(listing_urls=["https://allevents.in/{city}"]))

    assert summary.new_events == 1
    run = db_session.query(AgentRun).filter(AgentRun.agent_type == "discovery").one()
    assert run.status == "completed"
    assert "captcha" in run.error


# --- Excel export -------------------------------------------------------------------------


def _admin_user() -> User:
    return User(id=1, name="Admin", email="admin@100times.test", password_hash="x", role="ADMIN")


@pytest.fixture()
def client(db_session):
    app = FastAPI()
    app.include_router(discovery_routes.router, prefix="/api")
    app.dependency_overrides[get_db] = lambda: db_session
    app.dependency_overrides[get_current_user] = _admin_user
    return TestClient(app)


def _seed_events(db_session):
    organizer = DiscoveredOrganizer(name="Insignia Exhibitions", email="hello@insignia.example", phone="9876543210", website="https://insignia.example")
    db_session.add(organizer)
    db_session.flush()
    base = date(NEXT_YEAR, 11, 1)
    rows = [
        DiscoveredEvent(name="Jewellery World Exhibition", event_date=base, date=str(base), city="Mumbai", venue="JW Marriott Juhu",
                        url="https://allevents.in/mumbai/jwe", organizer_id=organizer.id, verification_status="VERIFIED", discovered_by="web_discovery_agent"),
        DiscoveredEvent(name="=HYPERLINK(\"http://evil.example\",\"click\")", event_date=base + timedelta(days=10), city="Pune",
                        verification_status="NEEDS_REVIEW", discovered_by="web_discovery_agent"),
        DiscoveredEvent(name="Outside Range Meetup", event_date=base + timedelta(days=60), city="Delhi", discovered_by="web_discovery_agent"),
        DiscoveredEvent(name="No Date Event", event_date=None, city="Delhi", discovered_by="web_discovery_agent"),
    ]
    db_session.add_all(rows)
    db_session.commit()


def _workbook(response):
    return load_workbook(io.BytesIO(response.content))


def test_export_filters_by_event_date_and_has_both_sheets(client, db_session):
    _seed_events(db_session)

    response = client.get(f"/api/discovery/export.xlsx?from={NEXT_YEAR}-11-01&to={NEXT_YEAR}-11-30")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert f"discovered-events_{NEXT_YEAR}-11-01_to_{NEXT_YEAR}-11-30.xlsx" in response.headers["content-disposition"]
    book = _workbook(response)
    assert book.sheetnames == ["Events", "Organizers"]
    events = list(book["Events"].iter_rows(values_only=True))
    header, data = events[0], events[1:]
    assert "Event name" in header and "Organizer name" in header
    names = [row[header.index("Event name")] for row in data]
    assert names == ["Jewellery World Exhibition", "=HYPERLINK(\"http://evil.example\",\"click\")"]
    organizers = list(book["Organizers"].iter_rows(values_only=True))
    assert len(organizers) == 2 and organizers[1][0] == "Insignia Exhibitions"


def test_export_never_writes_scraped_text_as_a_formula(client, db_session):
    _seed_events(db_session)

    book = _workbook(client.get(f"/api/discovery/export.xlsx?from={NEXT_YEAR}-11-01&to={NEXT_YEAR}-11-30"))

    for row in book["Events"].iter_rows():
        for cell in row:
            assert cell.data_type != "f"


def test_export_without_a_range_includes_everything(client, db_session):
    _seed_events(db_session)

    book = _workbook(client.get("/api/discovery/export.xlsx"))

    assert book["Events"].max_row == 5  # header + 4 events, undated ones included


def test_export_rejects_a_reversed_range(client):
    assert client.get(f"/api/discovery/export.xlsx?from={NEXT_YEAR}-12-01&to={NEXT_YEAR}-11-01").status_code == 400


def test_export_requires_admin():
    app = FastAPI()
    app.include_router(discovery_routes.router, prefix="/api")
    with TestClient(app) as anonymous:
        assert anonymous.get("/api/discovery/export.xlsx").status_code == 401


@pytest.mark.parametrize("title", [
    "Upcoming Trade Fairs & Exhibitions Worldwide 2026",
    "Technology Conferences in Bangalore 2026",
    "All Hackathons 2026",
])
def test_best_effort_ignores_real_listing_titles(title):
    html = f"<html><title>{title}</title><p>12 December {NEXT_YEAR}</p></html>"

    assert text_extractor.extract_best_effort(_source(), html, cities=[]) is None


def test_best_effort_ignores_breadcrumb_as_venue():
    html = f"<html><h1>Some Summit</h1><p>Location: Home</p><p>12 December {NEXT_YEAR}</p></html>"

    event = text_extractor.extract_best_effort(_source(), html, cities=[])

    assert event is not None and event.venue is None


def test_structured_data_html_entities_are_decoded(monkeypatch):
    url = "https://allevents.in/bangalore/parade/1"
    node = {"@type": "Event", "name": "Chaos &amp; Joy!", "url": url, "startDate": f"{NEXT_YEAR}-11-01",
            "location": {"name": "Rock &amp; Roll Hall", "address": {"addressLocality": "Bengaluru"}}}
    _serve(monkeypatch, {url: "<html><head>" + _ld(node) + "</head></html>"})

    event = page_extractor.extract_events_from_page(url, max_events=5)[0]

    assert event.name == "Chaos & Joy!"
    assert event.venue == "Rock & Roll Hall"


def test_run_categorizes_every_stored_event(db_session, monkeypatch):
    listing = "https://allevents.in/bangalore"
    nodes = [
        {"@type": "Event", "name": "Global Pharma Innovations Summit", "url": "https://allevents.in/bangalore/pharma/1", "startDate": f"{NEXT_YEAR}-10-16",
         "location": {"name": "Palace Grounds", "address": {"addressLocality": "Bengaluru"}}},
        {"@type": "Event", "name": "5G Networks & Telecommunications Meetup", "url": "https://allevents.in/bangalore/telecom/2", "startDate": f"{NEXT_YEAR}-10-24",
         "location": {"name": "EPAM", "address": {"addressLocality": "Bengaluru"}}},
    ]
    _serve(monkeypatch, {listing: "<html><head>" + "".join(_ld(n) for n in nodes) + "</head></html>"})
    _no_organizer(monkeypatch)

    orchestrator.run_discovery_agent(db_session, _config(listing_urls=["https://allevents.in/{city}"], search_depth=0))

    stored = {e.name: (e.category, e.subcategory) for e in db_session.query(DiscoveredEvent).all()}
    assert stored == {
        "Global Pharma Innovations Summit": ("Pharmaceuticals", "Summit"),
        "5G Networks & Telecommunications Meetup": ("Telecommunications", "Meetup"),
    }


@pytest.mark.parametrize("url, title", [
    ("https://allevents.in/pages/event-manager-app", "Manage your events with ease"),
    ("https://www.knowafest.com/events/page/2?country=India", "Explore events."),
    ("https://exhibition.org.in/exhibition-in-delhi/upcoming/all/india", "Plan Your Exhibition"),
    ("https://example.in/some-event", "Plan Your Exhibition"),
    ("https://example.in/events/upcoming", "Robotics Summit"),
])
def test_best_effort_skips_marketing_and_listing_pages(url, title):
    html = f"<html><h1>{title}</h1><p>25 October {NEXT_YEAR}</p></html>"

    assert text_extractor.extract_best_effort(_source(url), html, cities=[]) is None


def test_best_effort_keeps_real_titles_that_start_with_a_verb():
    html = f"<html><h1>Discover the Manchester MBA Information Session</h1><p>15 October {NEXT_YEAR}</p></html>"

    assert text_extractor.extract_best_effort(_source("https://example.in/mba-session"), html, cities=[]) is not None
